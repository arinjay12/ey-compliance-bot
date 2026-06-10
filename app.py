"""
app.py
------
Full Streamlit chat UI for the EY SEBI Compliance Bot.

Features (current):
  - Upload PDFs with duplicate detection (won't re-ingest same file twice)
  - Sidebar shows exactly which documents are in the knowledge base
  - Chat interface with history that persists after closing the app
  - Streaming responses (answer types out word by word)
  - Trailing prompt suggestions after each answer (Level 1)
  - Source excerpts from actual circulars
  - Level 2 monitor panel — shows last check time, conflicts found, run button
  - Level 3: Download chat history as PDF or Excel report

Usage:
    streamlit run app.py
"""

import os
import re
import json
import shutil
import hashlib
import tempfile
from io import BytesIO
from pathlib import Path
from datetime import datetime

import streamlit as st
from langchain_community.document_loaders import PyPDFLoader
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_community.vectorstores import Chroma
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

import rag_core   # shared RAG pipeline (retrieval, prompting, LLM, refusal gate)
import charts     # Level 4: turn numeric answers into charts
import voice      # Level 5: local speech-to-text (ask by voice)

# ── Config ─────────────────────────────────────────────────────────────────────
CHROMA_DIR     = "./chroma_db"
DOCS_DIR       = "./docs"
EMBED_MODEL    = "all-MiniLM-L6-v2"
LLM_MODEL      = "llama3"
HISTORY_FILE   = "./chat_history.json"    # persists chat across sessions
INGESTED_FILE  = "./ingested_docs.json"   # tracks which PDFs are already ingested
L2_STATUS_FILE = "./level2_status.json"   # written by level2.py, read here
AUDIT_FILE     = "./audit_log.jsonl"      # append-only compliance audit trail

os.makedirs(DOCS_DIR, exist_ok=True)


def ollama_up(timeout: float = 2.0) -> bool:
    """Quick check that the local Llama 3 server is reachable."""
    import urllib.request
    try:
        urllib.request.urlopen("http://localhost:11434/api/tags", timeout=timeout)
        return True
    except Exception:
        return False


def write_audit(question: str, answer: str, sources: list, refused: bool):
    """Append an audit record so every answer is traceable to its sources —
    a compliance tool should be able to show what it said and on what basis."""
    try:
        entry = {
            "timestamp": datetime.now().isoformat(timespec="seconds"),
            "question": question,
            "answer": answer,
            "refused": refused,
            "sources": [
                {"document": Path(d.metadata.get("source", "")).name,
                 "page": d.metadata.get("page")}
                for d in sources
            ],
        }
        with open(AUDIT_FILE, "a", encoding="utf-8") as f:
            f.write(json.dumps(entry, ensure_ascii=False) + "\n")
    except Exception:
        pass   # auditing must never break the chat

# ── Page setup ─────────────────────────────────────────────────────────────────
st.set_page_config(
    page_title="EY SEBI Compliance Bot",
    page_icon="📋",
    layout="wide",
)

st.markdown("""
<style>
    .header-bar {
        background: #ffe600;
        padding: 0.9rem 1.4rem;
        border-radius: 8px;
        margin-bottom: 1.2rem;
    }
    .header-bar h1 { color: #2e2e38; margin: 0; font-size: 1.55rem; font-weight: 700; }
    .header-bar p  { color: #2e2e38; margin: 0.2rem 0 0; font-size: 0.82rem; }
    .sug-label { color: #888; font-size: 0.8rem; font-weight: 600;
                 text-transform: uppercase; letter-spacing: 0.04em;
                 margin: 0.6rem 0 0.3rem; }
    .doc-chip { background: #ffe600; color: #2e2e38; border-radius: 4px;
                padding: 3px 8px; font-size: 0.75rem; margin: 3px 0;
                display: block; font-weight: 500; }
</style>
""", unsafe_allow_html=True)

st.markdown("""
<div class="header-bar">
    <h1>📋 EY SEBI Compliance Assistant</h1>
    <p>Upload SEBI circulars · Ask compliance queries · Powered by Llama 3 (local, no API cost)</p>
</div>
""", unsafe_allow_html=True)

# ── Cached resources ───────────────────────────────────────────────────────────
# The embedding model, vector store and LLM all live in rag_core (one source of
# truth, shared with the eval harness and Level 2). Wrap them so Streamlit shows
# a spinner on first load.
@st.cache_resource(show_spinner="Loading embedding model…")
def load_embeddings():
    return rag_core.get_embeddings()

@st.cache_resource(show_spinner="Connecting to Llama 3 via Ollama…")
def load_llm():
    return rag_core.get_llm()

# ── Ingested docs tracker ──────────────────────────────────────────────────────
def load_ingested() -> list:
    if os.path.exists(INGESTED_FILE):
        with open(INGESTED_FILE, "r") as f:
            return json.load(f)
    return []

def save_ingested(docs: list):
    with open(INGESTED_FILE, "w") as f:
        json.dump(docs, f)

# ── Chat history persistence ───────────────────────────────────────────────────
def load_history() -> list:
    if os.path.exists(HISTORY_FILE):
        with open(HISTORY_FILE, "r") as f:
            return json.load(f)
    return []

def save_history(messages: list):
    with open(HISTORY_FILE, "w") as f:
        json.dump(messages, f, ensure_ascii=False)

# ── Level 2 status reader ──────────────────────────────────────────────────────
def load_l2_status() -> dict:
    if os.path.exists(L2_STATUS_FILE):
        with open(L2_STATUS_FILE, "r") as f:
            return json.load(f)
    return {}

# ── Core functions ─────────────────────────────────────────────────────────────
def ingest_pdf(file_path: str, original_name: str, embeddings) -> tuple:
    """
    Load a PDF, split into chunks, store in ChromaDB.
    Returns (chunk_count, status) where status is one of:
      "ok" | "duplicate" | "no_text".
    "no_text" means almost no extractable text was found — usually a scanned /
    image-only PDF that would need OCR. We do NOT ingest it (so it isn't silently
    added as empty) and let the caller warn the user.
    """
    ingested = load_ingested()
    if original_name in ingested:
        return 0, "duplicate"

    pages       = PyPDFLoader(file_path).load()
    total_chars = sum(len((p.page_content or "").strip()) for p in pages)
    n_pages     = max(len(pages), 1)

    # Scanned/image PDFs extract little-to-no text. Flag if effectively empty
    # or under ~40 characters per page on average (a real circular has far more).
    if total_chars < 40 or (total_chars / n_pages) < 40:
        return 0, "no_text"

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=rag_core.CHUNK_SIZE if hasattr(rag_core, "CHUNK_SIZE") else 1000,
        chunk_overlap=150,
    )
    chunks   = splitter.split_documents(pages)
    if not chunks:
        return 0, "no_text"

    # Attach the same document-level metadata the batch ingester uses, so
    # uploaded PDFs behave identically (reference number in source tag, etc.).
    meta = rag_core.doc_meta_for(original_name, pages[0].page_content if pages else "")
    label = f"{meta['title']} | Ref: {meta['ref']} | Dated: {meta['date']}"
    for ch in chunks:
        ch.page_content = f"[{label}]\n{ch.page_content}"
        ch.metadata["doc_title"] = meta["title"]
        ch.metadata["doc_ref"]   = meta["ref"]
        ch.metadata["doc_date"]  = meta["date"]

    if os.path.exists(CHROMA_DIR):
        vs = Chroma(persist_directory=CHROMA_DIR, embedding_function=embeddings)
        vs.add_documents(chunks)
    else:
        Chroma.from_documents(chunks, embeddings, persist_directory=CHROMA_DIR)
    rag_core.reset_vectorstore()   # drop cached handle so new chunks are visible

    ingested.append(original_name)
    save_ingested(ingested)
    return len(chunks), "ok"


# Retrieval, prompting and the out-of-scope refusal gate all live in rag_core —
# the same code the evaluation harness and Level 2 use. The chat handler below
# calls rag_core.retrieve / build_prompt directly (passing chat history).


def generate_suggestions(query: str, answer: str, llm) -> list:
    """Generate 3 follow-up questions after each answer."""
    prompt = (
        f'A SEBI compliance analyst asked: "{query}"\n'
        f'The answer was: "{answer[:500]}"\n\n'
        "Suggest exactly 3 concise follow-up questions they might ask next about "
        "SEBI regulations or compliance obligations. "
        "Each must be a proper question ending with ?. "
        "Return only the 3 questions, one per line, no numbering or bullets."
    )
    try:
        raw   = llm.invoke(prompt)
        lines = [ln.strip() for ln in raw.strip().split("\n") if ln.strip()]
        return [ln for ln in lines if "?" in ln][:3]
    except Exception:
        return []


# ── Level 3: Export helpers ────────────────────────────────────────────────────

def _extract_qa_pairs(messages: list) -> list:
    """Convert flat message list into list of {question, answer} dicts."""
    pairs = []
    pending_q = None
    for msg in messages:
        if msg["role"] == "user":
            pending_q = msg["content"]
        elif msg["role"] == "assistant" and pending_q is not None:
            pairs.append({"question": pending_q, "answer": msg["content"]})
            pending_q = None
    return pairs


def generate_pdf_report(messages: list) -> bytes:
    """Build a formatted PDF compliance report from chat history."""
    buf    = BytesIO()
    doc    = SimpleDocTemplate(
        buf, pagesize=A4,
        leftMargin=2*cm, rightMargin=2*cm,
        topMargin=2*cm,  bottomMargin=2*cm,
    )
    styles = getSampleStyleSheet()

    # Custom styles
    title_style = ParagraphStyle(
        "EYTitle",
        parent=styles["Title"],
        fontSize=18, textColor=colors.HexColor("#2e2e38"),
        spaceAfter=4,
    )
    sub_style = ParagraphStyle(
        "EYSub",
        parent=styles["Normal"],
        fontSize=9, textColor=colors.HexColor("#666666"),
        spaceAfter=12,
    )
    q_style = ParagraphStyle(
        "Question",
        parent=styles["Heading3"],
        fontSize=11, textColor=colors.HexColor("#2e2e38"),
        spaceBefore=14, spaceAfter=4,
        leftIndent=0,
    )
    a_style = ParagraphStyle(
        "Answer",
        parent=styles["BodyText"],
        fontSize=10, leading=14,
        textColor=colors.HexColor("#1a1a1a"),
        leftIndent=12, spaceAfter=6,
    )

    story = []

    # Header bar (simulated with coloured paragraph)
    story.append(Paragraph("EY SEBI Compliance Assistant", title_style))
    story.append(Paragraph(
        f"Compliance Q&amp;A Report  ·  Generated {datetime.now().strftime('%d %B %Y, %H:%M')}",
        sub_style,
    ))
    story.append(HRFlowable(width="100%", thickness=2,
                             color=colors.HexColor("#ffe600"), spaceAfter=16))

    pairs = _extract_qa_pairs(messages)
    if not pairs:
        story.append(Paragraph("No conversation history to export.", styles["Normal"]))
    else:
        for i, qa in enumerate(pairs, 1):
            # Escape XML special chars for ReportLab
            q_text = qa["question"].replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            a_text = qa["answer"].replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            story.append(Paragraph(f"Q{i}. {q_text}", q_style))
            story.append(Paragraph(a_text, a_style))
            if i < len(pairs):
                story.append(HRFlowable(width="100%", thickness=0.5,
                                         color=colors.HexColor("#dddddd"), spaceAfter=4))

    story.append(Spacer(1, 1*cm))
    story.append(Paragraph(
        "This report was generated by the EY SEBI Compliance Bot (Llama 3 · ChromaDB · local).",
        sub_style,
    ))

    doc.build(story)
    return buf.getvalue()


def generate_excel_report(messages: list) -> bytes:
    """Build an Excel workbook with Q&A log from chat history."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Compliance Q&A"

    # Header row
    headers   = ["#", "Question", "Answer"]
    hdr_fill  = PatternFill("solid", fgColor="FFE600")   # EY yellow
    hdr_font  = Font(bold=True, color="2E2E38")

    for col, hdr in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=hdr)
        cell.fill   = hdr_fill
        cell.font   = hdr_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 55
    ws.column_dimensions["C"].width = 90
    ws.row_dimensions[1].height     = 20

    # Data rows
    pairs = _extract_qa_pairs(messages)
    for row_idx, qa in enumerate(pairs, 2):
        ws.cell(row=row_idx, column=1, value=row_idx - 1)
        ws.cell(row=row_idx, column=2, value=qa["question"])
        a_cell = ws.cell(row=row_idx, column=3, value=qa["answer"])
        a_cell.alignment = Alignment(wrap_text=True, vertical="top")

        # Zebra striping
        if row_idx % 2 == 0:
            fill = PatternFill("solid", fgColor="F8F8F8")
            for c in range(1, 4):
                ws.cell(row=row_idx, column=c).fill = fill

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── Session state init ─────────────────────────────────────────────────────────
if "messages" not in st.session_state:
    st.session_state.messages    = load_history()   # load from file on startup
if "suggestions" not in st.session_state:
    st.session_state.suggestions = []

# ── SIDEBAR ────────────────────────────────────────────────────────────────────
with st.sidebar:

    # ── 1. Document Upload ─────────────────────────────────────────────────────
    st.header("📁 Document Management")

    uploaded_files = st.file_uploader(
        "Upload SEBI / RBI Circulars (PDF)",
        type="pdf",
        accept_multiple_files=True,
    )

    embeddings = load_embeddings()

    if uploaded_files:
        if st.button("⚙️ Process Documents", type="primary", use_container_width=True):
            for f in uploaded_files:
                with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as tmp:
                    tmp.write(f.read())
                    tmp_path = tmp.name
                with st.spinner(f"Processing {f.name}…"):
                    try:
                        n, status = ingest_pdf(tmp_path, f.name, embeddings)
                        if status == "duplicate":
                            st.warning(f"⚠️ {f.name} — already ingested, skipped")
                        elif status == "no_text":
                            st.error(
                                f"🖼️ {f.name} — no readable text found. This looks "
                                "like a scanned / image-only PDF. It was NOT added "
                                "(the bot can only read text). OCR the file first, "
                                "then re-upload."
                            )
                        else:
                            st.success(f"✅ {f.name} — {n} chunks added")
                    except Exception as e:
                        st.error(f"❌ {f.name}: {e}")
                    finally:
                        os.unlink(tmp_path)

    st.divider()

    # ── 2. Documents in Knowledge Base ────────────────────────────────────────
    st.markdown("**📚 Knowledge Base**")
    ingested = load_ingested()
    if ingested:
        for doc_name in ingested:
            st.markdown(
                f'<span class="doc-chip">📄 {doc_name}</span>',
                unsafe_allow_html=True
            )
    else:
        st.caption("No documents loaded yet.")

    st.divider()

    # Clear everything button
    if st.button("🗑️ Clear Knowledge Base", use_container_width=True):
        if os.path.exists(CHROMA_DIR):
            shutil.rmtree(CHROMA_DIR)
        for f in [INGESTED_FILE, HISTORY_FILE]:
            if os.path.exists(f):
                os.remove(f)
        st.session_state.messages    = []
        st.session_state.suggestions = []
        st.success("Cleared.")
        st.rerun()

    # Clear chat only button
    if st.button("🧹 Clear Chat Only", use_container_width=True):
        st.session_state.messages    = []
        st.session_state.suggestions = []
        if os.path.exists(HISTORY_FILE):
            os.remove(HISTORY_FILE)
        st.rerun()

    st.divider()

    # ── 3. Level 3 — Export Chat ──────────────────────────────────────────────
    st.markdown("**📥 Export Chat History**")
    export_msgs = st.session_state.get("messages", [])
    has_chat    = any(m["role"] == "assistant" for m in export_msgs)

    if has_chat:
        col_pdf, col_xls = st.columns(2)

        with col_pdf:
            try:
                pdf_bytes = generate_pdf_report(export_msgs)
                fname_pdf = f"ey_compliance_{datetime.now().strftime('%Y%m%d_%H%M')}.pdf"
                st.download_button(
                    label="📄 PDF",
                    data=pdf_bytes,
                    file_name=fname_pdf,
                    mime="application/pdf",
                    use_container_width=True,
                )
            except Exception as e:
                st.caption(f"PDF error: {e}")

        with col_xls:
            try:
                xls_bytes = generate_excel_report(export_msgs)
                fname_xls = f"ey_compliance_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
                st.download_button(
                    label="📊 Excel",
                    data=xls_bytes,
                    file_name=fname_xls,
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    use_container_width=True,
                )
            except Exception as e:
                st.caption(f"Excel error: {e}")
    else:
        st.caption("Ask a question first to enable export.")

    st.divider()

    # ── 4. Level 2 Monitor Panel ───────────────────────────────────────────────
    st.header("🔍 Level 2 Monitor")
    st.caption("Auto-checks SEBI RSS feed daily for conflicts with stored docs")

    l2 = load_l2_status()

    if l2:
        st.caption(f"🕐 Last run: {l2.get('last_run', '—')}")
        st.caption(f"📋 Circulars checked: {l2.get('circulars_checked', 0)}")
        alerts = l2.get("alerts_found", 0)
        if alerts > 0:
            st.error(f"⚠️ {alerts} conflict(s) found — check your email")
            if l2.get("last_alert_title"):
                st.caption(f"Latest: {l2['last_alert_title'][:60]}…")
        else:
            st.success("✅ No conflicts detected")
    else:
        st.info("No checks run yet.\nRun level2.py to start.")

    if st.button("🔄 Run Check Now", use_container_width=True):
        with st.spinner("Fetching SEBI RSS and checking for conflicts… (may take a minute)"):
            try:
                import level2
                level2.run_check()
                st.success("Check complete!")
            except Exception as e:
                st.error(f"Error: {e}")
        st.rerun()

    st.divider()
    st.caption("EY Compliance Bot  ·  Llama 3 via Ollama  ·  ChromaDB")

# ── Chat history display ───────────────────────────────────────────────────────
for i, msg in enumerate(st.session_state.messages):
    with st.chat_message(msg["role"]):
        st.markdown(msg["content"])
        if msg.get("chart"):   # Level 4: re-render the chart with the answer
            st.caption("📊 Visualised from the answer")
            st.plotly_chart(charts.build_figure(msg["chart"]),
                            use_container_width=True, key=f"hist_chart_{i}")
        # Level 5 (output): read a real answer aloud, offline, on demand
        if msg["role"] == "assistant" and msg["content"] != rag_core.REFUSAL_LINE:
            if st.button("🔊 Listen", key=f"tts_{i}"):
                with st.spinner("Generating audio…"):
                    audio_bytes = voice.synthesize(msg["content"])
                if audio_bytes:
                    st.audio(audio_bytes, format="audio/wav", autoplay=True)
                else:
                    st.caption("Audio unavailable.")

# ── Trailing suggestions ───────────────────────────────────────────────────────
if st.session_state.suggestions:
    st.markdown('<p class="sug-label">💡 Suggested follow-up questions</p>',
                unsafe_allow_html=True)
    cols = st.columns(len(st.session_state.suggestions))
    for i, sug in enumerate(st.session_state.suggestions):
        if cols[i].button(sug, key=f"sug_{i}", use_container_width=True):
            st.session_state["pending"] = sug
            st.rerun()

# ── Voice input (Level 5) — ask by speaking ────────────────────────────────────
with st.expander("🎤 Ask by voice"):
    mic = st.audio_input("Record your question, then wait a moment for transcription")
    if mic is not None:
        sig = hashlib.md5(mic.getvalue()).hexdigest()
        if st.session_state.get("last_audio_sig") != sig:   # only handle new audio
            st.session_state["last_audio_sig"] = sig
            with st.spinner("Transcribing your question…"):
                spoken = voice.transcribe(mic.getvalue())
            if spoken:
                st.session_state["pending"] = spoken
                st.rerun()
            else:
                st.warning("Couldn't catch that — please try recording again.")

# ── Chat input ────────────────────────────────────────────────────────────────
user_input = None
if "pending" in st.session_state:
    user_input = st.session_state.pop("pending")
else:
    user_input = st.chat_input(
        "Ask about SEBI regulations, deadlines, UPI compliance, LODR obligations…"
    )

# ── Process query ─────────────────────────────────────────────────────────────
if user_input:
    if not os.path.exists(CHROMA_DIR):
        st.error("📂 Please upload and process at least one document first.")
        st.stop()
    if not ollama_up():
        st.error("🦙 Llama 3 isn't running. Start Ollama (`ollama serve` or the "
                 "Ollama app), then ask again.")
        st.stop()

    st.session_state.messages.append({"role": "user", "content": user_input})
    with st.chat_message("user"):
        st.markdown(user_input)

    with st.chat_message("assistant"):

        # 1. For a follow-up, fold in the previous question so retrieval has the
        #    topic context ("does that apply to commercial paper?" alone finds
        #    nothing). Self-contained questions are unchanged.
        history = st.session_state.messages[:-1]   # everything before this turn
        search_query = rag_core.build_search_query(user_input, history)

        # 2. Relevance gate — refuse out-of-scope queries instead of guessing
        if rag_core.min_distance(search_query) > rag_core.REFUSE_DISTANCE:
            answer  = rag_core.REFUSAL_LINE
            sources = []
            st.markdown(answer)
        else:
            # 3. Retrieve on the expanded query; the answer prompt keeps the
            #    natural question + recent conversation for reference resolution.
            with st.spinner("Searching circulars…"):
                sources = rag_core.retrieve(search_query)
                prompt  = rag_core.build_prompt(sources, user_input, history=history)

            # 4. Stream the LLM response word by word
            llm    = load_llm()
            answer = st.write_stream(llm.stream(prompt))   # streams & returns full string

        # 4. Show source excerpts + follow-up suggestions (only for real answers)
        chart_spec = None
        if sources:
            with st.expander("📄 Source excerpts from circulars"):
                for i, doc in enumerate(sources[:3], 1):
                    src  = Path(doc.metadata.get("source", "Unknown")).name
                    page = doc.metadata.get("page", "N/A")
                    st.caption(f"**[{i}]  {src}  —  Page {page}**")
                    st.text(doc.page_content[:280] + "…")
                    if i < min(3, len(sources)):
                        st.divider()

            # Level 4 — if the answer has chartable numbers, build a chart
            if charts.looks_chartable(answer):
                with st.spinner("Charting the numbers…"):
                    chart_spec = charts.extract_chart_data(
                        user_input, answer, load_llm()
                    )

            with st.spinner("Generating suggestions…"):
                st.session_state.suggestions = generate_suggestions(
                    user_input, answer, load_llm()
                )
        else:
            st.session_state.suggestions = []

    # 5. Audit trail — record what was answered and from which sources
    write_audit(user_input, answer, sources,
                refused=(answer == rag_core.REFUSAL_LINE))

    # 6. Save to session + persist to disk (chart spec rides along with the message)
    msg = {"role": "assistant", "content": answer}
    if chart_spec:
        msg["chart"] = chart_spec
    st.session_state.messages.append(msg)
    save_history(st.session_state.messages)
    st.rerun()
